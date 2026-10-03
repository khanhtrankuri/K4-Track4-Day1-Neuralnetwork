"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import openpyxl

# các cột công thức có sẵn trong template -> không bao giờ ghi đè
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

_LOSS_DISPLAY = {"ce": "CE", "mse": "MSE"}
_OPTIMIZER_DISPLAY = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


class _JSONEncoder(json.JSONEncoder):
    """Cho phép json.dump xử lý tuple/numpy scalar xuất hiện trong cfg/history/summary."""

    def default(self, o):
        if isinstance(o, (np.floating, np.integer)):
            return o.item()
        if isinstance(o, np.ndarray):
            return o.tolist()
        return super().default(o)


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    path = out_dir / f"{exp_id}.json"
    payload = {"cfg": result["cfg"], "history": result["history"], "summary": result["summary"]}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, cls=_JSONEncoder, ensure_ascii=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    paths = sorted(Path(results_dir).glob("*.json"), key=lambda p: p.stem)
    results = []
    for p in paths:
        with open(p, encoding="utf-8") as f:
            results.append(json.load(f))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.

    eval_scores: dict trả về bởi scripts/evaluate.py (--out ...json), tức có khoá "accuracy" và
    "macro_f1". Chỉ truyền cho baseline và cấu hình cuối cùng.
    """
    cfg, summary = result["cfg"], result["summary"]
    hidden = cfg["hidden"]
    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg["group"],
        "description": cfg.get("description", ""),
        "loss": _LOSS_DISPLAY.get(cfg["loss"], cfg["loss"]),
        "optimizer": _OPTIMIZER_DISPLAY.get(cfg["optimizer"], cfg["optimizer"]),
        "lr": cfg["lr"],
        "weight_decay": cfg["weight_decay"],
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": "-".join(str(h) for h in hidden),
        "dropout": cfg["dropout"],
        "clip_norm": "none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
        "step0_loss": summary["step0_loss"],
        "best_val_loss": summary["best_val_loss"],
        "best_epoch": summary["best_epoch"],
        "final_train_loss": summary["final_train_loss"],
        "final_val_loss": summary["final_val_loss"],
        "val_acc": summary["val_acc"],
        "val_macro_f1": summary["val_macro_f1"],
        "time_per_epoch_s": summary["time_per_epoch_s"],
        "peak_mem_MB": summary["peak_mem_MB"],
        "diverged": "Y" if summary["diverged"] else "N",
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes,
    }
    if eval_scores is not None:
        row["eval_acc"] = eval_scores["accuracy"]
        row["eval_macro_f1"] = eval_scores["macro_f1"]
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    wb = openpyxl.load_workbook(template_path)  # KHÔNG data_only=True, giữ nguyên công thức
    ws = wb["Experiments"]

    header = [cell.value for cell in ws[1]]
    col_of = {name: idx + 1 for idx, name in enumerate(header) if name and name not in FORMULA_COLUMNS}

    unknown = set()
    for r, row in enumerate(rows, start=2):
        for key, value in row.items():
            col = col_of.get(key)
            if col is None:
                unknown.add(key)
                continue
            ws.cell(row=r, column=col, value=value)

    if unknown:
        print(f"[write_xlsx] cảnh báo: bỏ qua các khoá không khớp cột nào: {sorted(unknown)}")

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
